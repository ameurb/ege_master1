import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Optional

from app.config import DATA_DIR, TP_FILES_DIR, SUBMISSIONS_DIR
from app.database import get_db

# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------


def get_settings() -> dict:
    with get_db() as cur:
        cur.execute("SELECT key, value FROM settings")
        rows = cur.fetchall()
    return {row["key"]: row["value"] for row in rows}


def save_settings(settings: dict) -> None:
    with get_db() as cur:
        for key, value in settings.items():
            cur.execute(
                "INSERT INTO settings (key, value) VALUES (%s, %s) "
                "ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value",
                (key, json.dumps(value)),
            )


def get_max_team_size() -> int:
    with get_db() as cur:
        cur.execute("SELECT value FROM settings WHERE key = 'max_team_size'")
        row = cur.fetchone()
    if row:
        return int(row["value"]) if not isinstance(row["value"], int) else row["value"]
    return 2


# ---------------------------------------------------------------------------
# Students
# ---------------------------------------------------------------------------


def get_students() -> list[dict]:
    with get_db() as cur:
        cur.execute("SELECT matricule, name, grp, team_id FROM students ORDER BY name")
        rows = cur.fetchall()
    return [{"matricule": r["matricule"], "name": r["name"], "group": r["grp"], "team_id": r["team_id"]} for r in rows]


def get_student(matricule: str) -> Optional[dict]:
    with get_db() as cur:
        cur.execute("SELECT matricule, name, grp, team_id FROM students WHERE matricule = %s", (matricule,))
        r = cur.fetchone()
    if r:
        return {"matricule": r["matricule"], "name": r["name"], "group": r["grp"], "team_id": r["team_id"]}
    return None


def add_student(matricule: str, name: str, group: str, binome_matricule: str = "") -> dict:
    student = {"matricule": matricule, "name": name, "group": group, "team_id": ""}
    with get_db() as cur:
        try:
            cur.execute(
                "INSERT INTO students (matricule, name, grp, team_id) VALUES (%s, %s, %s, %s)",
                (matricule, name, group, ""),
            )
        except Exception:
            raise ValueError(f"Student with matricule {matricule} already exists")
    return student


def set_team(matricules: list[str]) -> bool:
    if len(matricules) < 1 or len(matricules) > 3:
        return False

    team_key = make_team_key(matricules)

    with get_db() as cur:
        # Validate all exist
        for m in matricules:
            cur.execute("SELECT matricule FROM students WHERE matricule = %s", (m,))
            if not cur.fetchone():
                return False

        # Clear old teams: find all students who share a team_id with any of these students
        for m in matricules:
            cur.execute("SELECT team_id FROM students WHERE matricule = %s", (m,))
            row = cur.fetchone()
            old_team = row["team_id"] if row else ""
            if old_team:
                cur.execute("UPDATE students SET team_id = '' WHERE team_id = %s", (old_team,))

        # Set new team
        for m in matricules:
            cur.execute("UPDATE students SET team_id = %s WHERE matricule = %s", (team_key, m))

    return True


def get_team_members(matricule: str) -> list[dict]:
    student = get_student(matricule)
    if not student or not student.get("team_id"):
        return [student] if student else []
    team_id = student["team_id"]
    with get_db() as cur:
        cur.execute("SELECT matricule, name, grp, team_id FROM students WHERE team_id = %s ORDER BY name", (team_id,))
        rows = cur.fetchall()
    return [{"matricule": r["matricule"], "name": r["name"], "group": r["grp"], "team_id": r["team_id"]} for r in rows]


def make_team_key(matricules: list[str]) -> str:
    return "_".join(sorted(matricules))


def delete_student(matricule: str) -> bool:
    with get_db() as cur:
        cur.execute("DELETE FROM students WHERE matricule = %s", (matricule,))
        return cur.rowcount > 0


def import_students_csv(csv_text: str) -> dict:
    added = 0
    skipped = 0
    with get_db() as cur:
        for line in csv_text.strip().splitlines():
            line = line.strip()
            if not line or line.startswith("matricule"):
                continue
            parts = [p.strip() for p in line.split(",")]
            if len(parts) < 3:
                skipped += 1
                continue
            matricule, name, group = parts[0], parts[1], parts[2]
            try:
                cur.execute(
                    "INSERT INTO students (matricule, name, grp, team_id) VALUES (%s, %s, %s, '') "
                    "ON CONFLICT (matricule) DO NOTHING",
                    (matricule, name, group),
                )
                if cur.rowcount > 0:
                    added += 1
                else:
                    skipped += 1
            except Exception:
                skipped += 1
    return {"added": added, "skipped": skipped}


def get_students_by_group() -> dict[str, list[dict]]:
    groups: dict[str, list[dict]] = {}
    for s in get_students():
        g = s.get("group", "Unknown")
        groups.setdefault(g, []).append(s)
    return groups


# ---------------------------------------------------------------------------
# TPs
# ---------------------------------------------------------------------------


def get_tps() -> list[dict]:
    with get_db() as cur:
        cur.execute("SELECT id, title, description, week, due_date, created_at, files, steps FROM tps ORDER BY week, id")
        rows = cur.fetchall()
    result = []
    for r in rows:
        result.append({
            "id": r["id"],
            "title": r["title"],
            "description": r["description"],
            "week": r["week"],
            "due_date": r["due_date"],
            "created_at": r["created_at"].isoformat() if r["created_at"] else "",
            "files": r["files"] if isinstance(r["files"], list) else json.loads(r["files"]) if r["files"] else [],
            "steps": r["steps"] if isinstance(r["steps"], list) else json.loads(r["steps"]) if r["steps"] else [],
        })
    return result


def get_tp(tp_id: str) -> Optional[dict]:
    with get_db() as cur:
        cur.execute("SELECT id, title, description, week, due_date, created_at, files, steps FROM tps WHERE id = %s", (tp_id,))
        r = cur.fetchone()
    if r:
        return {
            "id": r["id"],
            "title": r["title"],
            "description": r["description"],
            "week": r["week"],
            "due_date": r["due_date"],
            "created_at": r["created_at"].isoformat() if r["created_at"] else "",
            "files": r["files"] if isinstance(r["files"], list) else json.loads(r["files"]) if r["files"] else [],
            "steps": r["steps"] if isinstance(r["steps"], list) else json.loads(r["steps"]) if r["steps"] else [],
        }
    return None


def create_tp(tp_id: str, title: str, description: str, week: int, due_date: str) -> dict:
    now = datetime.now()
    with get_db() as cur:
        try:
            cur.execute(
                "INSERT INTO tps (id, title, description, week, due_date, created_at, files, steps) "
                "VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
                (tp_id, title, description, week, due_date, now, json.dumps([]), json.dumps([])),
            )
        except Exception:
            raise ValueError(f"TP with id {tp_id} already exists")

    # Create TP files directory
    (TP_FILES_DIR / tp_id).mkdir(exist_ok=True)

    return {
        "id": tp_id,
        "title": title,
        "description": description,
        "week": week,
        "due_date": due_date,
        "created_at": now.isoformat(),
        "files": [],
        "steps": [],
    }


def update_tp(tp_id: str, title: str, description: str, week: int, due_date: str) -> Optional[dict]:
    with get_db() as cur:
        cur.execute(
            "UPDATE tps SET title=%s, description=%s, week=%s, due_date=%s WHERE id=%s",
            (title, description, week, due_date, tp_id),
        )
        if cur.rowcount == 0:
            return None
    return get_tp(tp_id)


def delete_tp(tp_id: str) -> bool:
    tp = get_tp(tp_id)
    if not tp:
        return False
    with get_db() as cur:
        cur.execute("DELETE FROM submissions WHERE tp_id = %s", (tp_id,))
        cur.execute("DELETE FROM grades WHERE tp_id = %s", (tp_id,))
        cur.execute("DELETE FROM tps WHERE id = %s", (tp_id,))
    tp_files_dir = TP_FILES_DIR / tp_id
    if tp_files_dir.exists():
        shutil.rmtree(tp_files_dir)
    sub_dir = SUBMISSIONS_DIR / tp_id
    if sub_dir.exists():
        shutil.rmtree(sub_dir)
    return True


def delete_tp_file(tp_id: str, filename: str) -> bool:
    """Remove a single attachment from a TP (DB list + disk)."""
    tp = get_tp(tp_id)
    if not tp:
        return False
    from app.services.file_service import sanitize_filename
    safe_name = sanitize_filename(filename)
    file_path = TP_FILES_DIR / tp_id / safe_name
    if file_path.exists():
        file_path.unlink()
    with get_db() as cur:
        cur.execute("SELECT files FROM tps WHERE id = %s", (tp_id,))
        row = cur.fetchone()
        if row:
            files = row["files"] if isinstance(row["files"], list) else json.loads(row["files"]) if row["files"] else []
            files = [f for f in files if f != safe_name]
            cur.execute("UPDATE tps SET files = %s WHERE id = %s", (json.dumps(files), tp_id))
    return True


def add_tp_file(tp_id: str, filename: str) -> None:
    with get_db() as cur:
        cur.execute("SELECT files FROM tps WHERE id = %s", (tp_id,))
        row = cur.fetchone()
        if row:
            files = row["files"] if isinstance(row["files"], list) else json.loads(row["files"]) if row["files"] else []
            if filename not in files:
                files.append(filename)
                cur.execute("UPDATE tps SET files = %s WHERE id = %s", (json.dumps(files), tp_id))


def toggle_tp_step(tp_id: str, step_num: int) -> Optional[dict]:
    with get_db() as cur:
        cur.execute("SELECT steps FROM tps WHERE id = %s", (tp_id,))
        row = cur.fetchone()
        if not row:
            return None
        steps = row["steps"] if isinstance(row["steps"], list) else json.loads(row["steps"]) if row["steps"] else []
        for step in steps:
            if step["num"] == step_num:
                step["enabled"] = not step["enabled"]
                cur.execute("UPDATE tps SET steps = %s WHERE id = %s", (json.dumps(steps), tp_id))
                return step
    return None


def get_tp_files_dir(tp_id: str) -> Path:
    d = TP_FILES_DIR / tp_id
    d.mkdir(exist_ok=True)
    return d


# ---------------------------------------------------------------------------
# Sessions (PC locking)
# ---------------------------------------------------------------------------


def get_active_session() -> Optional[dict]:
    with get_db() as cur:
        cur.execute(
            "SELECT id, tp_id, started_at, started_by, status, closed_at, pc_locks "
            "FROM sessions WHERE status = 'open' ORDER BY id DESC LIMIT 1"
        )
        r = cur.fetchone()
    if r:
        return {
            "tp_id": r["tp_id"],
            "started_at": r["started_at"].isoformat() if r["started_at"] else "",
            "started_by": r["started_by"],
            "status": r["status"],
            "pc_locks": r["pc_locks"] if isinstance(r["pc_locks"], dict) else json.loads(r["pc_locks"]) if r["pc_locks"] else {},
            "_id": r["id"],
        }
    return None


def start_session(tp_id: str = "all") -> dict:
    # Check if there's already an active session
    existing = get_active_session()
    if existing:
        raise ValueError("A session is already active. Close it first.")

    now = datetime.now()
    with get_db() as cur:
        cur.execute(
            "INSERT INTO sessions (tp_id, started_at, started_by, status, pc_locks) "
            "VALUES (%s, %s, 'professor', 'open', %s) RETURNING id",
            (tp_id, now, json.dumps({})),
        )
        row = cur.fetchone()

    return {
        "tp_id": tp_id,
        "started_at": now.isoformat(),
        "started_by": "professor",
        "status": "open",
        "pc_locks": {},
        "_id": row["id"],
    }


def stop_session() -> Optional[dict]:
    session = get_active_session()
    if not session:
        return None
    now = datetime.now()
    with get_db() as cur:
        cur.execute(
            "UPDATE sessions SET status = 'closed', closed_at = %s WHERE id = %s",
            (now, session["_id"]),
        )
    session["status"] = "closed"
    session["closed_at"] = now.isoformat()
    return session


def check_pc_lock(pc_id: str, team_key: str) -> str:
    # PC locking disabled — always allow
    return "ok"


def lock_pc(pc_id: str, team_key: str, team_names: str) -> None:
    session = get_active_session()
    if not session or session.get("status") != "open":
        return
    locks = session.get("pc_locks", {})
    locks[pc_id] = {
        "team_key": team_key,
        "team_names": team_names,
        "locked_at": datetime.now().isoformat(),
    }
    with get_db() as cur:
        cur.execute(
            "UPDATE sessions SET pc_locks = %s WHERE id = %s",
            (json.dumps(locks), session["_id"]),
        )


# ---------------------------------------------------------------------------
# Submissions
# ---------------------------------------------------------------------------


def make_binome_key(mat1: str, mat2: str) -> str:
    return "_".join(sorted([mat1, mat2]))


def get_team_key_for_student(matricule: str) -> str:
    student = get_student(matricule)
    if student and student.get("team_id"):
        return student["team_id"]
    return matricule


def get_submission_dir(tp_id: str, binome_key: str) -> Path:
    d = SUBMISSIONS_DIR / tp_id / binome_key
    d.mkdir(parents=True, exist_ok=True)
    return d


def save_submission_meta(tp_id: str, binome_key: str, meta: dict) -> None:
    # Save to database
    with get_db() as cur:
        cur.execute(
            "INSERT INTO submissions (tp_id, team_key, submitted_at, ip_address, pc_id, team_members, files) "
            "VALUES (%s, %s, %s, %s, %s, %s, %s) "
            "ON CONFLICT (tp_id, team_key) DO UPDATE SET "
            "submitted_at = EXCLUDED.submitted_at, ip_address = EXCLUDED.ip_address, "
            "pc_id = EXCLUDED.pc_id, team_members = EXCLUDED.team_members, files = EXCLUDED.files",
            (
                tp_id, binome_key,
                meta.get("submitted_at", datetime.now().isoformat()),
                meta.get("ip_address", ""),
                meta.get("pc_id", ""),
                json.dumps(meta.get("team_members", [])),
                json.dumps(meta.get("files", [])),
            ),
        )
    # Also keep the file-based meta for backwards compatibility with file serving
    d = get_submission_dir(tp_id, binome_key)
    d.parent.mkdir(parents=True, exist_ok=True)
    (d / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def get_submission_meta(tp_id: str, binome_key: str) -> Optional[dict]:
    with get_db() as cur:
        cur.execute(
            "SELECT tp_id, team_key, submitted_at, ip_address, pc_id, team_members, files "
            "FROM submissions WHERE tp_id = %s AND team_key = %s",
            (tp_id, binome_key),
        )
        r = cur.fetchone()
    if r:
        return {
            "tp_id": r["tp_id"],
            "team_key": r["team_key"],
            "submitted_at": r["submitted_at"].isoformat() if r["submitted_at"] else "",
            "ip_address": r["ip_address"],
            "pc_id": r["pc_id"],
            "team_members": r["team_members"] if isinstance(r["team_members"], list) else json.loads(r["team_members"]) if r["team_members"] else [],
            "files": r["files"] if isinstance(r["files"], list) else json.loads(r["files"]) if r["files"] else [],
        }
    return None


def list_submissions(tp_id: str) -> list[dict]:
    with get_db() as cur:
        cur.execute(
            "SELECT tp_id, team_key, submitted_at, ip_address, pc_id, team_members, files "
            "FROM submissions WHERE tp_id = %s ORDER BY submitted_at DESC",
            (tp_id,),
        )
        rows = cur.fetchall()
    submissions = []
    for r in rows:
        meta = {
            "tp_id": r["tp_id"],
            "team_key": r["team_key"],
            "binome_key": r["team_key"],
            "submitted_at": r["submitted_at"].isoformat() if r["submitted_at"] else "",
            "ip_address": r["ip_address"],
            "pc_id": r["pc_id"],
            "team_members": r["team_members"] if isinstance(r["team_members"], list) else json.loads(r["team_members"]) if r["team_members"] else [],
            "files": r["files"] if isinstance(r["files"], list) else json.loads(r["files"]) if r["files"] else [],
        }
        # Also check actual files on disk
        sub_dir = SUBMISSIONS_DIR / tp_id / r["team_key"]
        if sub_dir.exists():
            disk_files = [f.name for f in sub_dir.iterdir() if f.is_file() and f.name != "meta.json"]
            if disk_files:
                meta["files"] = disk_files
        submissions.append(meta)
    return submissions


def get_submission_files(tp_id: str, binome_key: str) -> list[dict]:
    d = SUBMISSIONS_DIR / tp_id / binome_key
    if not d.exists():
        return []
    files = []
    for f in sorted(d.iterdir()):
        if f.is_file() and f.name != "meta.json":
            files.append({
                "name": f.name,
                "size": f.stat().st_size,
                "extension": f.suffix.lower(),
            })
    return files


def get_submission_file_path(tp_id: str, binome_key: str, filename: str) -> Optional[Path]:
    p = SUBMISSIONS_DIR / tp_id / binome_key / filename
    if p.exists() and p.is_file():
        return p
    return None


# ---------------------------------------------------------------------------
# Attendance
# ---------------------------------------------------------------------------


def _ensure_notes_column():
    """Add notes column to attendance_sessions if it doesn't exist."""
    with get_db() as cur:
        cur.execute(
            "ALTER TABLE attendance_sessions ADD COLUMN IF NOT EXISTS notes JSONB DEFAULT '{}'"
        )


def _parse_notes(r) -> dict:
    notes = r.get("notes") if r else None
    if notes is None:
        return {}
    if isinstance(notes, dict):
        return notes
    return json.loads(notes) if notes else {}


def get_attendance() -> list[dict]:
    _ensure_notes_column()
    with get_db() as cur:
        cur.execute("SELECT date, week, label, records, notes FROM attendance_sessions ORDER BY date")
        rows = cur.fetchall()
    result = []
    for r in rows:
        result.append({
            "date": r["date"],
            "week": r["week"],
            "label": r["label"],
            "records": r["records"] if isinstance(r["records"], dict) else json.loads(r["records"]) if r["records"] else {},
            "notes": _parse_notes(r),
        })
    return result


def get_attendance_for_date(date: str) -> Optional[dict]:
    _ensure_notes_column()
    with get_db() as cur:
        cur.execute("SELECT date, week, label, records, notes FROM attendance_sessions WHERE date = %s", (date,))
        r = cur.fetchone()
    if r:
        return {
            "date": r["date"],
            "week": r["week"],
            "label": r["label"],
            "records": r["records"] if isinstance(r["records"], dict) else json.loads(r["records"]) if r["records"] else {},
            "notes": _parse_notes(r),
        }
    return None


def mark_students_present(matricules: list[str], session: dict) -> None:
    """Mark a list of students as 'present' in today's attendance, preserving existing records."""
    if not matricules or not session:
        return
    started_at = session.get("started_at", "")
    date_str = started_at[:10] if started_at else datetime.now().strftime("%Y-%m-%d")
    tp = get_tp(session.get("tp_id", ""))
    week = tp["week"] if tp else 0

    existing = get_attendance_for_date(date_str)
    if existing:
        records = existing["records"]
        notes = existing.get("notes", {})
        label = existing["label"]
        week = existing["week"] or week
    else:
        records = {}
        notes = {}
        label = f"TP Session {date_str}"

    for mat in matricules:
        records[mat] = "present"

    save_attendance_session(date_str, week, label, records, notes)


def auto_mark_attendance_on_session_stop(session: dict) -> None:
    """When a session stops, mark submitters as present and non-submitters as absent."""
    if not session:
        return
    started_at = session.get("started_at", "")
    date_str = started_at[:10] if started_at else datetime.now().strftime("%Y-%m-%d")
    tp_info = get_tp(session.get("tp_id", ""))
    week = tp_info["week"] if tp_info else 0

    # Collect all submitter matricules from submissions made during this session
    submitter_matricules: set[str] = set()
    tps = get_tps()
    for t in tps:
        subs = list_submissions(t["id"])
        for sub in subs:
            if sub.get("submitted_at", "") >= started_at:
                for member in sub.get("team_members", []):
                    submitter_matricules.add(member["matricule"])

    # Get all students
    all_students = get_students()
    all_matricules = {s["matricule"] for s in all_students}

    # Merge with existing attendance
    existing = get_attendance_for_date(date_str)
    if existing:
        records = existing["records"]
        notes = existing.get("notes", {})
        label = existing["label"]
        week = existing["week"] or week
    else:
        records = {}
        notes = {}
        label = f"TP Session {date_str}"

    # Mark submitters present, non-submitters absent (only if not already marked)
    for mat in submitter_matricules:
        records[mat] = "present"
    for mat in all_matricules:
        if mat not in records:
            records[mat] = "absent"

    save_attendance_session(date_str, week, label, records, notes)


def save_attendance_session(date: str, week: int, label: str, records: dict[str, str], notes: Optional[dict[str, float]] = None) -> dict:
    _ensure_notes_column()
    notes = notes or {}
    with get_db() as cur:
        cur.execute(
            "INSERT INTO attendance_sessions (date, week, label, records, notes) VALUES (%s, %s, %s, %s, %s) "
            "ON CONFLICT (date) DO UPDATE SET week = EXCLUDED.week, label = EXCLUDED.label, "
            "records = EXCLUDED.records, notes = EXCLUDED.notes",
            (date, week, label, json.dumps(records), json.dumps(notes)),
        )
    return {"date": date, "week": week, "label": label, "records": records, "notes": notes}


# ---------------------------------------------------------------------------
# Grades
# ---------------------------------------------------------------------------


def get_grades() -> dict:
    with get_db() as cur:
        cur.execute("SELECT tp_id, team_key, grade, max_grade, comment FROM grades")
        rows = cur.fetchall()
    result: dict = {}
    for r in rows:
        tp_id = r["tp_id"]
        if tp_id not in result:
            result[tp_id] = {}
        result[tp_id][r["team_key"]] = {
            "grade": float(r["grade"]),
            "max_grade": float(r["max_grade"]),
            "comment": r["comment"],
        }
    return result


def get_grades_for_tp(tp_id: str) -> dict:
    with get_db() as cur:
        cur.execute("SELECT team_key, grade, max_grade, comment FROM grades WHERE tp_id = %s", (tp_id,))
        rows = cur.fetchall()
    result = {}
    for r in rows:
        result[r["team_key"]] = {
            "grade": float(r["grade"]),
            "max_grade": float(r["max_grade"]),
            "comment": r["comment"],
        }
    return result


def save_grades_for_tp(tp_id: str, grades_data: dict[str, dict]) -> None:
    with get_db() as cur:
        for team_key, g in grades_data.items():
            cur.execute(
                "INSERT INTO grades (tp_id, team_key, grade, max_grade, comment) "
                "VALUES (%s, %s, %s, %s, %s) "
                "ON CONFLICT (tp_id, team_key) DO UPDATE SET "
                "grade = EXCLUDED.grade, max_grade = EXCLUDED.max_grade, comment = EXCLUDED.comment",
                (tp_id, team_key, g.get("grade", 0), g.get("max_grade", 20), g.get("comment", "")),
            )


def get_participation_grades() -> dict[str, float]:
    """Return {matricule: participation_grade} for all students."""
    attendance_sessions = get_attendance()
    students = get_students()
    total_sessions = len(attendance_sessions)
    if total_sessions == 0:
        return {}

    max_note = 2.0
    result = {}
    for student in students:
        mat = student["matricule"]
        notes_sum = 0.0
        for sess in attendance_sessions:
            status = sess.get("records", {}).get(mat, "")
            note = sess.get("notes", {}).get(mat)
            if status == "present":
                notes_sum += float(note) if note is not None else max_note
            elif status == "late":
                notes_sum += float(note) if note is not None else (max_note * 0.5)
        result[mat] = round(notes_sum / (total_sessions * max_note) * 20, 2)
    return result


# ---------------------------------------------------------------------------
# Evaluation helpers
# ---------------------------------------------------------------------------


def get_evaluation_data() -> dict:
    students = get_students()
    tps = get_tps()
    grades = get_grades()
    attendance_sessions = get_attendance()

    rows = []
    for student in students:
        mat = student["matricule"]
        name = student["name"]
        group = student.get("group", "")

        # Use team_id for grade lookup
        team_key = student.get("team_id", "") or mat

        # TP grades
        tp_grades = {}
        grade_values = []
        for tp in tps:
            tp_id = tp["id"]
            tp_grade = grades.get(tp_id, {}).get(team_key)
            tp_grades[tp_id] = tp_grade
            if tp_grade:
                grade_values.append(tp_grade["grade"] / tp_grade["max_grade"] * 20)

        # Attendance & participation notes
        present_count = 0
        late_count = 0
        total_sessions = len(attendance_sessions)
        participation_notes_sum = 0.0
        max_note = 2.0  # max note per session
        for sess in attendance_sessions:
            status = sess.get("records", {}).get(mat, "")
            note = sess.get("notes", {}).get(mat)
            if status == "present":
                present_count += 1
                # Use per-session note if set, otherwise default to max
                participation_notes_sum += float(note) if note is not None else max_note
            elif status == "late":
                late_count += 1
                participation_notes_sum += float(note) if note is not None else (max_note * 0.5)

        att_pct = round(present_count / total_sessions * 100) if total_sessions > 0 else 0

        # Participation grade: sum of notes / max possible, scaled to /20
        if total_sessions > 0:
            max_possible = total_sessions * max_note
            participation_grade = round(participation_notes_sum / max_possible * 20, 2)
        else:
            participation_grade = None

        avg = round(sum(grade_values) / len(grade_values), 2) if grade_values else None

        rows.append({
            "matricule": mat,
            "name": name,
            "group": group,
            "tp_grades": tp_grades,
            "average": avg,
            "attendance_pct": att_pct,
            "present_count": present_count,
            "late_count": late_count,
            "total_sessions": total_sessions,
            "participation_grade": participation_grade,
        })

    return {"rows": rows, "tps": tps, "attendance_sessions": attendance_sessions}


def export_csv() -> str:
    eval_data = get_evaluation_data()
    tps = eval_data["tps"]
    rows = eval_data["rows"]

    headers = ["Matricule", "Name", "Group"]
    for tp in tps:
        headers.append(f"{tp['id']} ({tp['title']})")
    headers.extend(["Average", "Participation /20", "Attendance %"])

    lines = [",".join(headers)]
    for row in rows:
        parts = [row["matricule"], f'"{row["name"]}"', row["group"]]
        for tp in tps:
            g = row["tp_grades"].get(tp["id"])
            parts.append(f'{g["grade"]}/{g["max_grade"]}' if g else "")
        parts.append(str(row["average"]) if row["average"] is not None else "")
        parts.append(str(row["participation_grade"]) if row.get("participation_grade") is not None else "")
        parts.append(str(row["attendance_pct"]))
        lines.append(",".join(parts))

    return "\n".join(lines)


def export_xlsx() -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    eval_data = get_evaluation_data()
    tps = eval_data["tps"]
    rows = eval_data["rows"]
    attendance_sessions = eval_data["attendance_sessions"]

    wb = Workbook()

    # --- Sheet 1: Evaluation ---
    ws = wb.active
    ws.title = "Evaluation"

    headers = ["Matricule", "Name", "Group"]
    for tp in tps:
        headers.append(f"{tp['id']} ({tp['title']})")
    headers.extend(["Average", "Participation /20", "Attendance %"])

    # Header style
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    # Data rows
    green_font = Font(color="1D6F42")
    red_font = Font(color="C00000")

    for r_idx, row in enumerate(sorted(rows, key=lambda x: x["name"]), 2):
        ws.cell(row=r_idx, column=1, value=row["matricule"]).border = thin_border
        ws.cell(row=r_idx, column=2, value=row["name"]).border = thin_border
        ws.cell(row=r_idx, column=3, value=row["group"]).border = thin_border

        col = 4
        for tp in tps:
            g = row["tp_grades"].get(tp["id"])
            cell = ws.cell(row=r_idx, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if g:
                cell.value = g["grade"]
                cell.number_format = '0.##'
                cell.font = green_font if g["grade"] >= 10 else red_font
            col += 1

        # Average
        avg_cell = ws.cell(row=r_idx, column=col, value=row["average"])
        avg_cell.border = thin_border
        avg_cell.alignment = Alignment(horizontal="center")
        avg_cell.number_format = '0.00'
        if row["average"] is not None:
            avg_cell.font = Font(bold=True, color="1D6F42" if row["average"] >= 10 else "C00000")
        col += 1

        # Participation
        part_cell = ws.cell(row=r_idx, column=col, value=row.get("participation_grade"))
        part_cell.border = thin_border
        part_cell.alignment = Alignment(horizontal="center")
        part_cell.number_format = '0.00'
        if row.get("participation_grade") is not None:
            part_cell.font = green_font if row["participation_grade"] >= 10 else red_font
        col += 1

        # Attendance %
        att_cell = ws.cell(row=r_idx, column=col, value=row["attendance_pct"])
        att_cell.border = thin_border
        att_cell.alignment = Alignment(horizontal="center")
        att_cell.number_format = '0'
        col += 1

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 30)

    # Freeze header row + first 3 columns
    ws.freeze_panes = "D2"

    # --- Sheet 2: Attendance ---
    ws2 = wb.create_sheet("Attendance")

    att_headers = ["Matricule", "Name", "Group"]
    for sess in attendance_sessions:
        att_headers.append(f"{sess['label']}\n{sess['date']}")

    for col, h in enumerate(att_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    present_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    absent_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    late_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    for r_idx, student in enumerate(sorted(rows, key=lambda x: x["name"]), 2):
        ws2.cell(row=r_idx, column=1, value=student["matricule"]).border = thin_border
        ws2.cell(row=r_idx, column=2, value=student["name"]).border = thin_border
        ws2.cell(row=r_idx, column=3, value=student["group"]).border = thin_border

        for c_idx, sess in enumerate(attendance_sessions, 4):
            mat = student["matricule"]
            status = sess.get("records", {}).get(mat, "")
            note = sess.get("notes", {}).get(mat)
            cell = ws2.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if status == "present":
                cell.value = f"P ({note})" if note is not None else "P"
                cell.fill = present_fill
                cell.font = Font(bold=True, color="1D6F42")
            elif status == "absent":
                cell.value = "A"
                cell.fill = absent_fill
                cell.font = Font(bold=True, color="C00000")
            elif status == "late":
                cell.value = f"L ({note})" if note is not None else "L"
                cell.fill = late_fill
                cell.font = Font(bold=True, color="9C6500")
            else:
                cell.value = "-"

    for col in ws2.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value).split("\n")[0]))
        ws2.column_dimensions[col_letter].width = min(max_len + 4, 30)

    ws2.freeze_panes = "D2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
